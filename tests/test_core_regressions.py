"""Regression coverage for retained DOCX claims and strict class index integrity."""

from io import BytesIO
import unittest

from docx import Document

from core.document_parser import extract_docx
from core.model import ModelError, resolve_label_mapping
from core.preprocessing import build_claims


class DocxClaimBoundaryTests(unittest.TestCase):
    def extract_claim_texts(self, document):
        buffer = BytesIO()
        document.save(buffer)
        return [row["text"] for row in build_claims(extract_docx(buffer.getvalue()))]

    def test_separate_unpunctuated_paragraphs_remain_separate_claims(self):
        document = Document()
        document.add_paragraph("Emisi karbon berkurang sebesar 20% pada 2024")
        document.add_paragraph("Energi terbarukan mencapai 30% konsumsi energi")
        self.assertEqual(self.extract_claim_texts(document), [
            "Emisi karbon berkurang sebesar 20% pada 2024",
            "Energi terbarukan mencapai 30% konsumsi energi",
        ])

    def test_word_list_items_do_not_require_literal_bullet_characters(self):
        document = Document()
        claims = [
            "Emisi karbon berkurang sebesar 20% pada 2024",
            "Energi terbarukan mencapai 30% konsumsi energi",
        ]
        for claim in claims:
            document.add_paragraph(claim, style="List Bullet")
        self.assertEqual(self.extract_claim_texts(document), claims)

    def test_table_rows_remain_separate_and_keep_body_order(self):
        document = Document()
        document.add_paragraph("Pembukaan laporan lingkungan perusahaan")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Penurunan emisi karbon terhadap 2020"
        table.cell(0, 1).text = "20%"
        table.cell(1, 0).text = "Proporsi konsumsi energi terbarukan"
        table.cell(1, 1).text = "30%"
        document.add_paragraph("Penutup dan hasil verifikasi independen")
        self.assertEqual(self.extract_claim_texts(document), [
            "Pembukaan laporan lingkungan perusahaan",
            "Penurunan emisi karbon terhadap 2020 | 20%",
            "Proporsi konsumsi energi terbarukan | 30%",
            "Penutup dan hasil verifikasi independen",
        ])

    def test_soft_line_wrap_within_paragraph_still_joins(self):
        document = Document()
        document.add_paragraph("Penggunaan energi\nterbarukan meningkat sebesar 30% pada 2024")
        self.assertEqual(self.extract_claim_texts(document), [
            "Penggunaan energi terbarukan meningkat sebesar 30% pada 2024",
        ])


class StrictLabelIndexTests(unittest.TestCase):
    generic_config = {"id2label": {"0": "LABEL_0", "1": "LABEL_1"}}

    def test_semantic_metadata_rejects_boolean_fractional_and_coercible_indices(self):
        for invalid in (True, False, 0.9, 1.9, 0.0, 1.0, "0.9", "01", " 1", "+1", None, 2):
            with self.subTest(index=invalid), self.assertRaises(ModelError):
                resolve_label_mapping(self.generic_config, {"greenwashing_index": invalid})

    def test_valid_integer_or_exact_string_metadata_keeps_both_class_orders(self):
        for index in (0, 1, "0", "1"):
            with self.subTest(index=index):
                mapping = resolve_label_mapping(self.generic_config, {"greenwashing_index": index})
                self.assertEqual(mapping.greenwashing_index, int(index))
                self.assertEqual(mapping.low_indication_index, 1 - int(index))

    def test_inverse_mapping_cannot_mask_fractional_or_boolean_indices(self):
        for invalid in (0.9, False):
            config = {
                "id2label": {"0": "Lower Risk", "1": "Higher Risk"},
                "label2id": {"Lower Risk": invalid, "Higher Risk": 1},
            }
            with self.subTest(index=invalid), self.assertRaises(ModelError):
                resolve_label_mapping(config)

    def test_forward_mapping_rejects_noninteger_keys_with_no_inverse(self):
        for invalid in (0.9, False, "00"):
            with self.subTest(index=invalid), self.assertRaises(ModelError):
                resolve_label_mapping({"id2label": {invalid: "Lower Risk", 1: "Higher Risk"}})

    def test_valid_config_directions_are_unchanged(self):
        mapping = resolve_label_mapping({
            "id2label": {"0": "Lower Risk", "1": "Higher Risk"},
            "label2id": {"Lower Risk": 0, "Higher Risk": 1},
        })
        self.assertEqual(mapping.greenwashing_index, 1)
        self.assertEqual(mapping.id2label, {0: "Lower Risk", 1: "Higher Risk"})


if __name__ == "__main__":
    unittest.main()
