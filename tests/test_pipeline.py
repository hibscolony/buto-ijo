"""Deterministic parsing, segmentation, aggregation, and export unit tests.

Numeric probability fixtures below test only the aggregation mathematics. They
are never used as model predictions or exposed in the application.
"""

import csv
from io import BytesIO, StringIO
import json
import math
import unittest

import fitz
from docx import Document

from core.document_parser import (
    DocumentError,
    SCAN_PDF_MESSAGE,
    ScanPDFError,
    extract_docx,
    extract_pdf,
    extract_txt,
    parse_document,
)
from core.export import RESULT_COLUMNS, results_to_csv, safe_export_name, summary_to_json
from core.preprocessing import build_claims, clean_pages, clean_text, segment_claims
from core.risk import apply_threshold, calculate_document_risk


class DocumentParserTests(unittest.TestCase):
    def test_pdf_keeps_original_page_numbers_including_blank_pages(self):
        with fitz.open() as document:
            document.new_page().insert_text((72, 72), "Emisi karbon turun 25% pada tahun 2024.")
            document.new_page()
            document.new_page().insert_text((72, 72), "Energi terbarukan mencapai 30% konsumsi energi.")
            data = document.tobytes()
        pages = parse_document(data, "Sustainability_Report.PDF")
        self.assertEqual([page["page"] for page in pages], [1, 2, 3])
        self.assertIn("25%", pages[0]["text"])
        self.assertFalse(pages[1]["text"].strip())
        claims = build_claims(pages)
        self.assertEqual([claim["page"] for claim in claims], [1, 3])
        self.assertEqual([claim["claim_id"] for claim in claims], [1, 2])

    def test_pdf_without_text_layer_has_actionable_message(self):
        with fitz.open() as document:
            document.new_page().draw_rect(fitz.Rect(70, 70, 170, 170))
            data = document.tobytes()
        with self.assertRaises(ScanPDFError) as caught:
            extract_pdf(data)
        self.assertEqual(str(caught.exception), SCAN_PDF_MESSAGE)

    def test_encrypted_pdf_is_rejected(self):
        with fitz.open() as document:
            document.new_page().insert_text((72, 72), "Dokumen lingkungan rahasia.")
            data = document.tobytes(
                encryption=fitz.PDF_ENCRYPT_AES_256,
                owner_pw="test-owner-password",
                user_pw="test-reader-password",
            )
        with self.assertRaisesRegex(DocumentError, "terenkripsi|kata sandi"):
            extract_pdf(data)

    def test_empty_corrupt_and_unsupported_documents_are_rejected(self):
        for extractor in (extract_pdf, extract_docx, extract_txt):
            with self.subTest(extractor=extractor.__name__):
                with self.assertRaises(DocumentError):
                    extractor(b"")
        for extractor in (extract_pdf, extract_docx):
            with self.subTest(corrupt=extractor.__name__):
                with self.assertRaises(DocumentError):
                    extractor(b"This is not a valid document container")
        with self.assertRaisesRegex(DocumentError, "tidak didukung"):
            parse_document(b"some content", "report.xlsx")

    def test_docx_preserves_body_and_table_order(self):
        document = Document()
        document.add_paragraph("Pembukaan laporan lingkungan perusahaan.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Emisi karbon"
        table.cell(0, 1).text = "12,5%"
        document.add_paragraph("Penutup dan hasil verifikasi independen.")
        buffer = BytesIO()
        document.save(buffer)
        pages = extract_docx(buffer.getvalue())
        self.assertEqual(len(pages), 1)
        self.assertEqual(pages[0]["page"], 1)
        text = pages[0]["text"]
        self.assertLess(text.index("Pembukaan"), text.index("Emisi karbon"))
        self.assertLess(text.index("12,5%"), text.index("Penutup"))

    def test_txt_accepts_utf8_bom_and_preserves_unicode(self):
        content = "Emisi turun 12,5%; suhu dijaga pada 25°C."
        pages = extract_txt(content.encode("utf-8-sig"))
        self.assertEqual(pages, [{"page": 1, "text": content}])
        for invalid in (b"\xff\xfe", b"   \n\t", b"abc\x00def"):
            with self.subTest(invalid=invalid):
                with self.assertRaises(DocumentError):
                    extract_txt(invalid)


class PreprocessingTests(unittest.TestCase):
    def test_cleaning_retains_numbers_punctuation_units_and_case(self):
        source = "  Emisi\u00a0  CO2 turun 12,5%\t(1.200 kWh).\r\n\r\n\r\nTidak ada stemming! "
        self.assertEqual(
            clean_text(source),
            "Emisi CO2 turun 12,5% (1.200 kWh).\n\nTidak ada stemming!",
        )

    def test_sentence_boundaries_keep_abbreviations_and_decimal_points(self):
        source = (
            "PT. Hijau melaporkan penurunan emisi sebesar 12.5% pada 2024. "
            "Konsumsi energi tercatat sebesar 1.200 kWh pada fasilitas utama."
        )
        claims = segment_claims(source)
        self.assertEqual(len(claims), 2)
        self.assertTrue(claims[0].startswith("PT. Hijau"))
        self.assertIn("12.5%", claims[0])
        self.assertIn("1.200 kWh", claims[1])

    def test_short_fragments_and_whitespace_are_ignored(self):
        self.assertEqual(segment_claims(" \n\t "), [])
        self.assertEqual(segment_claims("Emisi turun."), [])
        self.assertEqual(segment_claims("Emisi turun.", min_char_length=5), ["Emisi turun."])
        for invalid in (0, -1, True):
            with self.subTest(invalid=invalid):
                with self.assertRaises(ValueError):
                    segment_claims("Emisi turun.", min_char_length=invalid)

    def test_wrapped_lines_and_bullets_create_complete_claims(self):
        source = (
            "Penggunaan energi\nterbarukan meningkat sebesar 30% pada tahun 2024.\n"
            "• Emisi karbon turun 20% dibandingkan tahun dasar 2020.\n"
            "• Pengambilan air berkurang 12% pada fasilitas yang sama."
        )
        claims = segment_claims(source)
        self.assertEqual(len(claims), 3)
        self.assertIn("energi terbarukan", claims[0])

    def test_repeated_headers_are_removed_without_removing_full_sentences(self):
        sentence = "Kami menggunakan energi terbarukan di seluruh fasilitas."
        pages = [
            {"page": index, "text": f"LAPORAN KEBERLANJUTAN\n{sentence}\nHalaman {index}"}
            for index in (1, 2, 3)
        ]
        cleaned = clean_pages(pages)
        self.assertEqual([page["text"] for page in cleaned], [sentence] * 3)
        one_line_pages = [{"page": index, "text": sentence} for index in (1, 2, 3)]
        self.assertEqual(clean_pages(one_line_pages), one_line_pages)


class RiskAndExportTests(unittest.TestCase):
    def setUp(self):
        # Deliberate mathematical inputs, not simulated model output.
        self.rows = [
            {"claim_id": 1, "page": 4, "claim": 'Emisi turun "12,5%", menurut audit.', "greenwashing_probability": 0.8},
            {"claim_id": 2, "page": 9, "claim": "Konsumsi energi terbarukan meningkat.", "greenwashing_probability": 0.2},
        ]

    def test_formula_uses_all_claims_as_ratio_denominator(self):
        summary = calculate_document_risk(self.rows)
        self.assertEqual(summary["total_claims"], 2)
        self.assertEqual(summary["flagged_claims"], 1)
        self.assertEqual(summary["flagged_ratio"], 0.5)
        self.assertEqual(summary["high_confidence_flagged_ratio"], 0.5)
        self.assertEqual(summary["mean_greenwashing_probability"], 0.5)
        self.assertAlmostEqual(summary["risk_score"], 50.0)
        self.assertEqual(summary["risk_level"], "MODERATE")

    def test_threshold_copies_rows_and_sets_selected_class_probability(self):
        first = apply_threshold(self.rows, threshold=0.9)
        self.assertTrue(all(row["prediction"] == "Low Indication" for row in first))
        self.assertAlmostEqual(first[0]["confidence"], 0.2)
        self.assertNotIn("prediction", self.rows[0])
        self.assertIsNot(first[0], self.rows[0])
        self.assertEqual(first[0]["greenwashing_probability"], 0.8)
        tie = apply_threshold(self.rows, threshold=0.8)
        self.assertEqual(tie[0]["prediction"], "Potential Greenwashing")
        self.assertAlmostEqual(calculate_document_risk(self.rows, threshold=0.9)["risk_score"], 25.0)

    def test_empty_and_nonfinite_probabilities_are_rejected(self):
        with self.assertRaises(ValueError):
            calculate_document_risk([])
        for value in (-0.1, 1.1, math.nan, math.inf):
            with self.subTest(value=value):
                with self.assertRaises(ValueError):
                    apply_threshold(self.rows, threshold=value)
                with self.assertRaises(ValueError):
                    calculate_document_risk([{"greenwashing_probability": value}])

    def test_risk_extremes_and_fractional_boundary_have_no_category_gaps(self):
        for probability, level in ((0.0, "LOW"), (1.0, "HIGH"), (0.6, "HIGH")):
            with self.subTest(probability=probability):
                summary = calculate_document_risk([{"greenwashing_probability": probability}])
                self.assertEqual(summary["risk_level"], level)
                self.assertGreaterEqual(summary["risk_score"], 0)
                self.assertLessEqual(summary["risk_score"], 100)
        # threshold=1 means p=.61 yields score 30.5, which must be MODERATE.
        fractional = calculate_document_risk([{"greenwashing_probability": 0.61}], threshold=1.0)
        self.assertAlmostEqual(fractional["risk_score"], 30.5)
        self.assertEqual(fractional["risk_level"], "MODERATE")

    def test_csv_round_trip_preserves_claims_and_actual_numbers(self):
        classified = apply_threshold(self.rows)
        data = results_to_csv(classified)
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        reader = csv.DictReader(StringIO(data.decode("utf-8-sig")))
        self.assertEqual(reader.fieldnames, RESULT_COLUMNS)
        exported = list(reader)
        self.assertEqual(exported[0]["claim"], self.rows[0]["claim"])
        self.assertEqual(float(exported[0]["greenwashing_probability"]), 0.8)
        self.assertEqual(exported[1]["page"], "9")
        self.assertEqual(len(exported), 2)

    def test_csv_formula_like_text_is_escaped(self):
        rows = [{**self.rows[0], "claim": '=HYPERLINK("https://example.invalid")'}]
        exported = list(csv.DictReader(StringIO(results_to_csv(rows).decode("utf-8-sig"))))
        self.assertTrue(exported[0]["claim"].startswith("'="))

    def test_json_summary_and_safe_filename(self):
        summary = calculate_document_risk(self.rows)
        document = "Laporan Keberlanjutan – 2024.pdf"
        decoded = json.loads(summary_to_json(summary, document).decode("utf-8"))
        self.assertEqual(decoded["document"], document)
        self.assertEqual(decoded["total_claims"], 2)
        self.assertEqual(decoded["risk_score"], summary["risk_score"])
        self.assertEqual(decoded["threshold"], 0.5)
        self.assertEqual(safe_export_name(r"..\..\Laporan 2024.pdf"), "Laporan 2024")
        self.assertNotIn("/", safe_export_name("../../report.pdf"))


if __name__ == "__main__":
    unittest.main()
